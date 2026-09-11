import re
import shutil
import sys
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

GENERAL_SHEET = "Generalidades"
DATA_SHEET = "Asistencia"

STUDENTS = [
    (1001, "10000000000001", "ALVARADO QUISPE, MARIA FERNANDA"),
    (1002, "10000000000002", "BAUTISTA HUAMAN, JOSE LUIS"),
    (1003, "10000000000003", "CHAVEZ MAMANI, ROSA ELENA"),
    (1004, "10000000000004", "DIAZ ROJAS, CARLOS ALBERTO"),
    (1005, "10000000000005", "ESPINOZA VARGAS, LUZ MARINA"),
]

DAYS_IN_MONTH = 30

HEADER_FILL = PatternFill("solid", fgColor="FFD9D9D9")
TITLE_FILL = PatternFill("solid", fgColor="FF1F4E79")
THIN = Side(style="thin", color="FF808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_generalidades(sheet):
    sheet["A1"] = "MINISTERIO DE EDUCACION - SIAGIE"
    sheet["A1"].font = Font(bold=True, size=12, color="FFFFFFFF")
    sheet["A1"].fill = TITLE_FILL
    sheet["A2"] = "REGISTRO DE ASISTENCIA MENSUAL - 6 PRIMARIA EBR"
    sheet["A2"].font = Font(bold=True, size=10)
    sheet["A3"] = "IE: 12345 - SETIEMBRE 2026"
    sheet["A3"].font = Font(bold=True, size=10)
    sheet["A5"] = "SECCION: A"
    sheet["A5"].font = Font(size=9)
    sheet["A6"] = "No editar esta hoja."
    sheet["A6"].font = Font(size=8, italic=True)
    sheet.column_dimensions["A"].width = 40


def build_asistencia(sheet):
    sheet["A1"] = "MINISTERIO DE EDUCACION - SIAGIE"
    sheet["A1"].font = Font(bold=True, size=12, color="FFFFFFFF")
    sheet["A1"].fill = TITLE_FILL
    sheet["A2"] = "REGISTRO DE ASISTENCIA MENSUAL - 6 PRIMARIA EBR"
    sheet["A2"].font = Font(bold=True, size=10)
    sheet["A3"] = "MES: SETIEMBRE 2026"
    sheet["A3"].font = Font(bold=True, size=10)

    headers = ["ID", "CodEstudiante", "Nombres"] + [str(day) for day in range(1, DAYS_IN_MONTH + 1)]
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=3, column=index, value=header)
        cell.font = Font(bold=True, size=9)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")

    first_data_row = 4
    for offset, (student_id, code, name) in enumerate(STUDENTS):
        row = first_data_row + offset
        sheet.cell(row=row, column=1, value=student_id).border = BORDER
        code_cell = sheet.cell(row=row, column=2, value=code)
        code_cell.border = BORDER
        code_cell.alignment = Alignment(horizontal="left")
        sheet.cell(row=row, column=3, value=name).border = BORDER
        for day in range(1, DAYS_IN_MONTH + 1):
            sheet.cell(row=row, column=3 + day).border = BORDER

    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 40
    for day in range(1, DAYS_IN_MONTH + 1):
        from openpyxl.utils import get_column_letter
        sheet.column_dimensions[get_column_letter(3 + day)].width = 4
    sheet.freeze_panes = "D4"
    sheet.sheet_view.showGridLines = False


INLINE_CELL = re.compile(r'<c r="([A-Z]+\d+)"((?: [a-z]+="[^"]*")*) t="inlineStr"><is><t>(.*?)</t></is></c>')

SHARED_STRINGS_PART = "xl/sharedStrings.xml"
SHARED_STRINGS_TYPE = (
    '<Override PartName="/xl/sharedStrings.xml" '
    'ContentType="application/vnd.openxmlformats-officedocument'
    '.spreadsheetml.sharedStrings+xml" />'
)
SHARED_STRINGS_REL = (
    '<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006'
    '/relationships/sharedStrings" Target="sharedStrings.xml" Id="rId6" />'
)


def to_shared_strings(destination):
    staging = destination + ".inline"
    shutil.move(destination, staging)
    table = []
    index_of = {}
    with zipfile.ZipFile(staging) as source:
        entries = [(item, source.read(item.filename)) for item in source.infolist()]

    def intern(text):
        if text not in index_of:
            index_of[text] = len(table)
            table.append(text)
        return index_of[text]

    def rewrite(match):
        return '<c r="%s"%s t="s"><v>%d</v></c>' % (
            match.group(1),
            match.group(2),
            intern(match.group(3)),
        )

    rewritten = []
    for item, payload in entries:
        if item.filename.startswith("xl/worksheets/sheet"):
            payload = INLINE_CELL.sub(rewrite, payload.decode("utf-8")).encode("utf-8")
        elif item.filename == "[Content_Types].xml":
            payload = payload.decode("utf-8").replace(
                "</Types>", SHARED_STRINGS_TYPE + "</Types>"
            ).encode("utf-8")
        elif item.filename == "xl/_rels/workbook.xml.rels":
            payload = payload.decode("utf-8").replace(
                "</Relationships>", SHARED_STRINGS_REL + "</Relationships>"
            ).encode("utf-8")
        rewritten.append((item.filename, payload))

    shared = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>']
    shared.append(
        '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'count="%d" uniqueCount="%d">' % (len(table), len(table))
    )
    for text in table:
        shared.append("<si><t>" + text.replace("&", "&amp;").replace("<", "&lt;") + "</t></si>")
    shared.append("</sst>")

    with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as target:
        for name, payload in rewritten:
            target.writestr(name, payload)
        target.writestr(SHARED_STRINGS_PART, "".join(shared))


def main(destination):
    workbook = Workbook()
    workbook.remove(workbook.active)
    build_generalidades(workbook.create_sheet(GENERAL_SHEET))
    build_asistencia(workbook.create_sheet(DATA_SHEET))
    workbook.save(destination)
    to_shared_strings(destination)


if __name__ == "__main__":
    main(sys.argv[1])
